import streamlit as st
from translate_codebase import (load_dictionaries, train_translation_model, translate_with_model,
                                translate_with_google, translate_word, translate_excel, 
                                translate_pptx, translate_srt, translate_csv, save_translated_file)
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
import io
import base64
import pandas as pd
import logging

# Configure logging
logging.basicConfig(level=logging.INFO)

def get_download_link(file_path, file_label):
    # Function to create a download link for a file
    with open(file_path, "rb") as f:
        bytes_data = f.read()
    b64 = base64.b64encode(bytes_data).decode()
    return f'<a href="data:file/txt;base64,{b64}" download="{file_path}">{file_label}</a>'

def get_csv_download_link(dataframe, file_label):
    # Function to create a download link for a CSV file
    csv = dataframe.to_csv(index=False).encode()
    b64 = base64.b64encode(csv).decode()
    href = f'<a href="data:file/csv;base64,{b64}" download="{file_label}.csv">Download {file_label} CSV</a>'
    return href

# Initialize translation dictionary and models
translation_dict = {}
models = None
ratings = []

st.title("XPLORE Translation")

# File uploader for CSV files to load dictionaries or train models
uploaded_csv_files = st.file_uploader("Upload CSV files for contextualized translation", type="csv", accept_multiple_files=True)
if uploaded_csv_files:
    if st.button("Load AI Model"):
        translation_dict = load_dictionaries(uploaded_csv_files)
        st.write("AI Model loaded successfully.")
    elif st.button("Train your AI model"):
        models = train_translation_model(uploaded_csv_files)
        st.write("Your AI model was trained successfully.")

st.markdown("---")

# Text area for entering text to translate
text_to_translate = st.text_area("Enter text to translate")
target_language = st.selectbox("Translate to", ["english", "indonesian", "malay", "hindi"])

if st.button("Translate Text"):
    # Translate the entered text
    if text_to_translate:
        if models:
            translated_text = translate_with_model(text_to_translate, 'auto', target_language, models)
        else:
            translated_text = translate_with_google(text_to_translate, target_language, translation_dict)
        st.write("Translated Text:")
        st.write(translated_text)
    else:
        st.error("Please enter text to translate.")

st.markdown("---")

# File uploader for document translations
uploaded_file = st.file_uploader("Upload a document", type=["docx", "xlsx", "pptx", "srt", "csv"])
if uploaded_file:
    file_type = uploaded_file.name.split('.')[-1]
    initial_name = uploaded_file.name.split('.')[0]
    if file_type == "docx":
        # Translate Word document
        doc = Document(uploaded_file)
        translated_doc = translate_word(doc, target_language, translation_dict, models)
        file_name = save_translated_file(translated_doc, file_type, initial_name, target_language)
        st.success(f"Translated Word document saved as '{file_name}'")
        download_link = get_download_link(file_name, 'Download Translated Document')
        st.markdown(download_link, unsafe_allow_html=True)
    elif file_type == "xlsx":
        # Translate Excel spreadsheet
        wb = load_workbook(uploaded_file)
        translated_wb = translate_excel(wb, target_language, translation_dict, models)
        file_name = save_translated_file(translated_wb, file_type, initial_name, target_language)
        st.success(f"Translated Excel spreadsheet saved as '{file_name}'")
        download_link = get_download_link(file_name, 'Download Translated Document')
        st.markdown(download_link, unsafe_allow_html=True)
    elif file_type == "pptx":
        # Translate PowerPoint presentation
        prs = Presentation(uploaded_file)
        translated_prs = translate_pptx(prs, target_language, translation_dict, models)
        file_name = save_translated_file(translated_prs, file_type, initial_name, target_language)
        st.success(f"Translated PowerPoint presentation saved as '{file_name}'")
        download_link = get_download_link(file_name, 'Download Translated Document')
        st.markdown(download_link, unsafe_allow_html=True)
    elif file_type == "srt":
        # Translate SRT subtitle file
        srt_content = uploaded_file.read().decode("utf-8")
        translated_srt = translate_srt(srt_content, target_language, translation_dict, models)
        file_name = save_translated_file(translated_srt, file_type, initial_name, target_language)
        st.success(f"Translated SRT file saved as '{file_name}'")
        download_link = get_download_link(file_name, 'Download Translated Document')
        st.markdown(download_link, unsafe_allow_html=True)
    elif file_type == "csv":
        # Translate CSV file
        df = pd.read_csv(uploaded_file)
        translated_df = translate_csv(df, target_language, translation_dict, models)
        file_name = save_translated_file(translated_df, file_type, initial_name, target_language)
        st.success(f"Translated CSV file saved as '{file_name}'")
        download_link = get_csv_download_link(translated_df, f'{initial_name}_{target_language}')
        st.markdown(download_link, unsafe_allow_html=True)
    else:
        st.error("Unsupported file type")

st.markdown("---")

# Rating section for translation accuracy
rating_options = {
    1: "1 - Not accurate",
    2: "2 - Less accurate",
    3: "3 - Quite accurate",
    4: "4 - Accurate",
    5: "5 - Very accurate"
}
accuracy_rating = st.selectbox("Rate the translation accuracy", options=list(rating_options.keys()), format_func=lambda x: rating_options[x], index=4)
st.write(f"Translation Accuracy Rating: {accuracy_rating} - {rating_options[accuracy_rating]}")

if st.button("Submit Rating"):
    # Save the rating
    ratings.append({"Rating": accuracy_rating, "Description": rating_options[accuracy_rating]})
    st.write("Rating saved successfully.")
